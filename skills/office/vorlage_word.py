"""Word-Vorlage im Hausdesign – das Gegenstueck zu ``vorlage.py`` (PowerPoint).

WARUM ES DAS GIBT: ``office_create_word`` rief ``Document()`` OHNE Argument
auf. Damit gilt das eingebaute Standarddokument von python-docx – und das ist
in drei Punkten falsch fuer ein Firmendokument:

  * **US-Letter** (21,59 × 27,94 cm) statt A4. Beim Drucken auf A4 fehlt unten
    ein Streifen, und Word zeigt beim Empfaenger „Papierformat passt nicht".
  * **Calibri/Cambria** und das Office-Standard-Blau in den Ueberschriften –
    kein Bezug zum Branding.
  * Keine Kopf- und Fusszeile: kein Logo, keine Seitenzahl.

ERGEBNIS: eine .docx mit echten Formatvorlagen (Standard, Titel, Ueberschrift
1–3, Aufzaehlung), deren Farben auf das THEME zeigen (``w:themeColor``) – wer
das Dokument in Word ueber „Entwurf → Farben" umstellt, bekommt ein stimmiges
Ergebnis, statt dass einzelne Absaetze ihre feste Farbe behalten. Das ist
dieselbe Ueberlegung wie bei PowerPoint: Theme aendern, nicht den Absatz.

WOHER DIE WERTE STAMMEN: Farben, Schriften und die Hausfarbe kommen aus
``vorlage.py`` (dort aus der Firmenvorlage ausgelesen) – EINE Quelle fuer
beide Formate, sonst laufen Folie und Dokument beim naechsten Farbwechsel
auseinander. Schriftgroessen und Raender sind die ueblichen Werte fuer
A4-Geschaeftsdokumente nach DIN 5008 (linker Rand 2,5 cm); eine
Firmen-Wordvorlage gibt es dafuer nicht – wer eine hat, laedt sie im Reiter
Branding hoch, dann gilt deren Gestaltung.

Fuer den PDF-Export gilt derselbe Vorbehalt wie bei PowerPoint (siehe
``vorlage.py``, ACHTUNG PDF-EXPORT): die CI-Schrift ist auf dem Server nicht
installiert.
"""

import io
import re
import zipfile
from pathlib import Path

from skills.office import vorlage as _pptx

VORLAGEN_DIR = _pptx.VORLAGEN_DIR
STANDARD_NAME = "standard.docx"

# ── Seite (Twips: 1 cm = 567) ───────────────────────────────────────────────
# A4 hochkant. Raender nach DIN 5008: links 2,5 cm (Lochung/Heftung), rechts
# 2 cm. Oben 2,7 cm, damit die Kopfzeile mit Logo (Abstand 1,25 cm) Luft zum
# Text hat.
A4_BREITE_CM = 21.0
A4_HOEHE_CM = 29.7
RAND_LINKS_CM = 2.5
RAND_RECHTS_CM = 2.0
RAND_OBEN_CM = 2.7
RAND_UNTEN_CM = 2.0
KOPF_ABSTAND_CM = 1.25
FUSS_ABSTAND_CM = 1.0

# Logo in der Kopfzeile: Hoehe fest, Breite aus dem Seitenverhaeltnis der
# Datei (python-docx rechnet sie selbst, wenn nur die Hoehe gesetzt ist).
LOGO_HOEHE_CM = 0.9

# ── Typografie (Punkt) ──────────────────────────────────────────────────────
# 10,5 pt Grundtext: die HelveticaNeue laeuft breit, 11 pt wirkt damit groeber
# als Calibri 11. Ueberschriften in der Titelschrift (majorFont).
TYPO_TEXT = 10.5
TYPO_TITEL = 24
TYPO_UNTERTITEL = 13
TYPO_H = [16, 13, 11]                 # Ueberschrift 1..3
TYPO_FUSS = 8
ZEILENABSTAND = 1.15
ABSTAND_NACH = 6                      # nach jedem Absatz
H_ABSTAND_VOR = [18, 12, 10]          # vor Ueberschrift 1..3
H_ABSTAND_NACH = [6, 4, 3]

# Farbrollen je Formatvorlage: (Theme-Slot, Hexwert). Der Hexwert ist der
# Rueckfall fuer Programme, die das Theme nicht auswerten (LibreOffice beim
# PDF-Export nimmt w:val, Word nimmt w:themeColor).
#   Titel + Ueberschrift 1 in der Hausfarbe, Ueberschrift 2 dunkel,
#   Ueberschrift 3 im Sekundaerblau der Palette – so bleibt die Hausfarbe
#   ein Akzent und wird nicht zur Flaeche.
ROLLE_TITEL = ("accent1", None)            # None = Branding-Akzent
ROLLE_H = [("accent1", None), ("text1", _pptx.DUNKEL), ("accent2", _pptx.PALETTE[0])]
ROLLE_GRAU = ("accent5", _pptx.GRAU)


def _w(tag: str) -> str:
    from docx.oxml.ns import qn
    return qn(f"w:{tag}")


def _farbe_setzen(font, rolle: tuple, akzent: str) -> None:
    """Setzt eine Theme-Farbe MIT festem Rueckfallwert an einen Zeichenstil."""
    _farbe_an_rpr(font.element.get_or_add_rPr(), rolle, akzent)


def _farbe_an_rpr(rpr, rolle: tuple, akzent: str) -> None:
    """Wie ``_farbe_setzen``, aber direkt an einem ``w:rPr``.

    ``get_or_add_color`` statt ``append``: das Schema legt die REIHENFOLGE der
    Kinder fest (``w:color`` vor ``w:sz``), und Word meldet bei falscher
    Reihenfolge „unlesbarer Inhalt" – LibreOffice dagegen nicht, ein
    Server-PDF zeigt den Fehler also nie."""
    slot, wert = rolle
    c = rpr.get_or_add_color()
    c.set(_w("val"), wert or akzent)
    c.set(_w("themeColor"), slot)


def _schrift(font, name: str) -> None:
    """Schrift fuer ALLE Schriftbereiche setzen.

    ``font.name`` setzt nur ascii/hAnsi; ohne eastAsia/cs nimmt Word fuer
    Sonderzeichen und Ziffern in manchen Sprachen wieder Calibri – das sieht
    man erst im fertigen Dokument als Schriftmischung."""
    font.name = name
    rf = font.element.get_or_add_rPr().get_or_add_rFonts()
    for attr in ("ascii", "hAnsi", "eastAsia", "cs"):
        rf.set(_w(attr), name)
    # Theme-Verweise entfernen, sonst gewinnen sie gegen den Namen.
    for attr in ("asciiTheme", "hAnsiTheme", "eastAsiaTheme", "cstheme"):
        if rf.get(_w(attr)) is not None:
            del rf.attrib[_w(attr)]


def _stile(doc, farben: dict) -> None:
    """Formatvorlagen auf das Hausdesign stellen."""
    from docx.shared import Pt
    akzent = farben["akzent"]
    st = doc.styles

    normal = st["Normal"]
    _schrift(normal.font, _pptx.SCHRIFT_TEXT)
    normal.font.size = Pt(TYPO_TEXT)
    _farbe_setzen(normal.font, ("text1", _pptx.DUNKEL), akzent)
    pf = normal.paragraph_format
    pf.line_spacing = ZEILENABSTAND
    pf.space_after = Pt(ABSTAND_NACH)

    titel = st["Title"]
    _schrift(titel.font, _pptx.SCHRIFT_TITEL)
    titel.font.size = Pt(TYPO_TITEL)
    titel.font.bold = True
    _farbe_setzen(titel.font, ROLLE_TITEL, akzent)
    titel.paragraph_format.space_after = Pt(12)
    # Die Standardvorlage zieht unter den Titel eine blaue Linie (pBdr) – in
    # der Hausfarbe waere sie eine zweite Akzentflaeche direkt unter der
    # ersten. Weg damit.
    ppr = titel.element.get_or_add_pPr()
    for b in ppr.findall(_w("pBdr")):
        ppr.remove(b)

    try:
        unter = st["Subtitle"]
        _schrift(unter.font, _pptx.SCHRIFT_TEXT)
        unter.font.size = Pt(TYPO_UNTERTITEL)
        unter.font.italic = False
        _farbe_setzen(unter.font, ROLLE_GRAU, akzent)
    except KeyError:
        pass

    for i in range(3):
        try:
            h = st[f"Heading {i + 1}"]
        except KeyError:
            continue
        _schrift(h.font, _pptx.SCHRIFT_TITEL)
        h.font.size = Pt(TYPO_H[i])
        h.font.bold = True
        h.font.italic = False
        _farbe_setzen(h.font, ROLLE_H[i], akzent)
        h.paragraph_format.space_before = Pt(H_ABSTAND_VOR[i])
        h.paragraph_format.space_after = Pt(H_ABSTAND_NACH[i])
        h.paragraph_format.keep_with_next = True


def _seite(doc) -> None:
    """A4 und Raender – fuer JEDEN Abschnitt, nicht nur den ersten."""
    from docx.shared import Cm
    for sec in doc.sections:
        sec.page_width = Cm(A4_BREITE_CM)
        sec.page_height = Cm(A4_HOEHE_CM)
        sec.left_margin = Cm(RAND_LINKS_CM)
        sec.right_margin = Cm(RAND_RECHTS_CM)
        sec.top_margin = Cm(RAND_OBEN_CM)
        sec.bottom_margin = Cm(RAND_UNTEN_CM)
        sec.header_distance = Cm(KOPF_ABSTAND_CM)
        sec.footer_distance = Cm(FUSS_ABSTAND_CM)


def _feld(absatz, code: str) -> None:
    """Haengt ein Word-Feld (PAGE, NUMPAGES) als einfaches Feld an.

    ``w:fldSimple`` statt der Dreiteilung begin/instr/end: es ist ein Element
    statt fuenf und wird von Word UND LibreOffice beim Oeffnen aktualisiert.
    Der Platzhaltertext '1' steht drin, damit eine Vorschau ohne
    Feldaktualisierung keine leere Stelle zeigt."""
    from docx.oxml import OxmlElement
    fld = OxmlElement("w:fldSimple")
    fld.set(_w("instr"), f" {code} ")
    r = OxmlElement("w:r")
    t = OxmlElement("w:t")
    t.text = "1"
    r.append(t)
    fld.append(r)
    absatz._p.append(fld)


def _unterstrich(absatz, farbe: str) -> None:
    """Duenne Linie unter einem Absatz (Kopfzeile) in der Hausfarbe."""
    from docx.oxml import OxmlElement
    ppr = absatz._p.get_or_add_pPr()
    bdr = OxmlElement("w:pBdr")
    unten = OxmlElement("w:bottom")
    unten.set(_w("val"), "single")
    unten.set(_w("sz"), "6")            # Achtelpunkt → 0,75 pt
    unten.set(_w("space"), "4")
    unten.set(_w("color"), farbe)
    unten.set(_w("themeColor"), "accent1")
    bdr.append(unten)
    # Schemareihenfolge: pBdr steht VOR shd/tabs/spacing/ind/jc. Ein
    # angehaengtes pBdr hinter jc rendert LibreOffice ohne Farbe – und Word
    # darf die Datei als beschaedigt ablehnen.
    ppr.insert_element_before(
        bdr, "w:shd", "w:tabs", "w:suppressAutoHyphens", "w:kinsoku",
        "w:wordWrap", "w:overflowPunct", "w:topLinePunct", "w:autoSpaceDE",
        "w:autoSpaceDN", "w:bidi", "w:adjustRightInd", "w:snapToGrid",
        "w:spacing", "w:ind", "w:contextualSpacing", "w:mirrorIndents",
        "w:suppressOverlap", "w:jc", "w:textDirection", "w:textAlignment",
        "w:textboxTightWrap", "w:outlineLvl", "w:divId", "w:cnfStyle",
        "w:rPr", "w:sectPr", "w:pPrChange")


def _kopf_und_fuss(doc, farben: dict) -> None:
    """Kopfzeile: Logo rechtsbuendig (sonst Firmenname), Linie in der Hausfarbe.
    Fusszeile: Firmenname links, „Seite X von Y" rechts."""
    from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
    from docx.shared import Cm, Pt
    firma = farben.get("firma") or ""
    logo = _pptx._logo_pfad()
    satzbreite = Cm(A4_BREITE_CM - RAND_LINKS_CM - RAND_RECHTS_CM)

    for sec in doc.sections:
        kopf = sec.header.paragraphs[0] if sec.header.paragraphs else sec.header.add_paragraph()
        kopf.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        gesetzt = False
        if logo:
            try:
                kopf.add_run().add_picture(str(logo), height=Cm(LOGO_HOEHE_CM))
                gesetzt = True
            except Exception:  # noqa: BLE001
                # SVG oder defektes Bild: python-docx kann es nicht einbetten.
                # Dann steht der Name da – eine Kopfzeile ohne Absender waere
                # schlechter als eine ohne Logo.
                gesetzt = False
        if not gesetzt and firma:
            r = kopf.add_run(firma)
            r.font.size = Pt(9)
            r.font.bold = True
            _farbe_setzen(r.font, ROLLE_GRAU, farben["akzent"])
        _unterstrich(kopf, farben["akzent"])

        fuss = sec.footer.paragraphs[0] if sec.footer.paragraphs else sec.footer.add_paragraph()
        fuss.paragraph_format.tab_stops.add_tab_stop(satzbreite, WD_TAB_ALIGNMENT.RIGHT)
        teile = [firma, "\t", "Seite "]
        for t in teile:
            if t:
                r = fuss.add_run(t)
                r.font.size = Pt(TYPO_FUSS)
                _farbe_setzen(r.font, ROLLE_GRAU, farben["akzent"])
        _feld(fuss, "PAGE")
        r = fuss.add_run(" von ")
        r.font.size = Pt(TYPO_FUSS)
        _farbe_setzen(r.font, ROLLE_GRAU, farben["akzent"])
        _feld(fuss, "NUMPAGES")
        # Die Feld-Laeufe erben sonst Groesse und Farbe des Absatzstils – die
        # Seitenzahl stuende in 10,5 pt Schwarz neben 8 pt Grau.
        for fld in fuss._p.findall(_w("fldSimple")):
            for run in fld.findall(_w("r")):
                rpr = run.get_or_add_rPr()
                _farbe_an_rpr(rpr, ROLLE_GRAU, farben["akzent"])
                rpr.get_or_add_sz().val = Pt(TYPO_FUSS)


def _patch_theme(docx_pfad: Path, farben: dict) -> None:
    """Farb- und Schriftschema im Theme ersetzen (wie bei PowerPoint).

    Dieselbe Funktion ``_theme_xml`` wie fuer die Folien – nur der Ort im ZIP
    unterscheidet sich (``word/theme/`` statt ``ppt/theme/``)."""
    tmp = _pptx.zwischenname(docx_pfad, ".tmp.docx")
    with zipfile.ZipFile(docx_pfad) as quelle:
        with zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as ziel:
            for eintrag in quelle.infolist():
                daten = quelle.read(eintrag.filename)
                if re.fullmatch(r"word/theme/theme\d+\.xml", eintrag.filename):
                    try:
                        daten = _pptx._theme_xml(daten, farben)
                    except Exception:  # noqa: BLE001
                        pass       # lieber unveraendertes Theme als kaputte Datei
                ziel.writestr(eintrag, daten)
    tmp.replace(docx_pfad)


def erzeuge(ziel: Path = None) -> Path:
    """Erzeugt die Hausvorlage aus den Branding-Werten und gibt ihren Pfad zurueck."""
    from docx import Document

    ziel = Path(ziel) if ziel else (VORLAGEN_DIR / STANDARD_NAME)
    ziel.parent.mkdir(parents=True, exist_ok=True)
    farben = _pptx.branding_farben()

    doc = Document()
    _seite(doc)
    _stile(doc, farben)
    _kopf_und_fuss(doc, farben)
    # Die Standardvorlage traegt „python-docx" als Autor und einen
    # Beispieltitel in den Eigenschaften – das erbte sonst jedes Dokument.
    cp = doc.core_properties
    cp.author = farben.get("firma") or ""
    cp.title = ""
    cp.subject = ""
    cp.keywords = ""
    cp.comments = ""
    cp.last_modified_by = ""

    # Eindeutiger, versteckter Zwischenname (siehe vorlage.ERZEUGE_SPERRE).
    tmp = _pptx.zwischenname(ziel, ".neu.docx")
    try:
        doc.save(str(tmp))
        # Theme erst NACH dem Speichern – python-docx schreibt den Theme-Part
        # unveraendert aus seiner Standardvorlage.
        _patch_theme(tmp, farben)
        tmp.replace(ziel)
    finally:
        if tmp.exists():
            tmp.unlink()
    return ziel


def sicherstellen(pfad: Path = None) -> Path:
    """Gibt die Hausvorlage zurueck und erzeugt sie beim ersten Mal.

    Keine Neuerzeugung bei jedem Aufruf: die Datei darf durch eine echte
    Firmenvorlage ersetzt sein (gleiche Regel wie bei PowerPoint)."""
    ziel = Path(pfad) if pfad else (VORLAGEN_DIR / STANDARD_NAME)
    if ziel.exists() and ziel.stat().st_size > 0:
        return ziel
    # Dieselbe Sperre wie bei PowerPoint – gleichzeitige Erzeugung (siehe dort).
    with _pptx.ERZEUGE_SPERRE:
        if ziel.exists() and ziel.stat().st_size > 0:
            return ziel
        return erzeuge(ziel)


def _ist_zwischenstand(name: str) -> bool:
    return name.startswith(".") or name.endswith((".tmp.docx", ".neu.docx"))


def verfuegbare() -> list:
    """Alle hinterlegten Word-Vorlagen (Dateinamen), Hausvorlage zuerst."""
    try:
        dateien = sorted(p.name for p in VORLAGEN_DIR.glob("*.docx")
                         if not _ist_zwischenstand(p.name))
    except Exception:  # noqa: BLE001
        return []
    if STANDARD_NAME in dateien:
        dateien.remove(STANDARD_NAME)
        dateien.insert(0, STANDARD_NAME)
    return dateien


def loese_vorlage(name: str = "") -> Path:
    """Loest einen Vorlagennamen auf; leer oder unbekannt = Hausvorlage.

    Nur der Dateiname zaehlt (der Name kommt aus einem LLM-Aufruf)."""
    if not name:
        return sicherstellen()
    sauber = Path(str(name)).name
    if not sauber.lower().endswith(".docx"):
        sauber += ".docx"
    kandidat = VORLAGEN_DIR / sauber
    if kandidat.exists() and not _ist_zwischenstand(sauber):
        return kandidat
    return sicherstellen()


# ── Hochgeladene Firmenvorlagen ─────────────────────────────────────────────
CT_DOKUMENT = "application/vnd.openxmlformats-officedocument.wordprocessingml.document.main+xml"
CT_VORLAGE = "application/vnd.openxmlformats-officedocument.wordprocessingml.template.main+xml"


def als_dokument(daten: bytes) -> bytes:
    """Macht aus einer .dotx eine .docx (nur der Inhaltstyp aendert sich).

    WARUM: python-docx oeffnet eine .dotx NICHT – ``Document()`` prueft den
    Inhaltstyp des Hauptteils und wirft „file is not a Word file, content
    type is '...template.main+xml'". Anders als bei PowerPoint genuegt das
    Umbenennen also nicht: die Vorlage laege ab, und JEDES Word-Dokument
    scheiterte danach. Inhaltlich sind beide Formate gleich; getauscht wird
    nur die eine Zeile in ``[Content_Types].xml``."""
    with zipfile.ZipFile(io.BytesIO(daten)) as quelle:
        ct = quelle.read("[Content_Types].xml")
        if CT_VORLAGE.encode() not in ct:
            return daten
        raus = io.BytesIO()
        with zipfile.ZipFile(raus, "w", zipfile.ZIP_DEFLATED) as ziel:
            for eintrag in quelle.infolist():
                inhalt = quelle.read(eintrag.filename)
                if eintrag.filename == "[Content_Types].xml":
                    inhalt = inhalt.replace(CT_VORLAGE.encode(), CT_DOKUMENT.encode())
                ziel.writestr(eintrag, inhalt)
    return raus.getvalue()


def inhalt_leeren(doc) -> int:
    """Entfernt den Textkoerper einer Vorlage, behaelt das Seitenformat.

    Eine Firmen-Wordvorlage bringt oft Mustertext mit („Titel hier eingeben",
    Beispieltabelle). Uebernommen werden Formatvorlagen, Kopf-/Fusszeilen und
    Seitenformat – die stehen NICHT im Koerper, sondern in eigenen Parts bzw.
    im abschliessenden ``w:sectPr``. Rueckgabe: Zahl der entfernten Elemente."""
    body = doc.element.body
    weg = 0
    for kind in list(body):
        if kind.tag == _w("sectPr"):
            continue
        body.remove(kind)
        weg += 1
    return weg
